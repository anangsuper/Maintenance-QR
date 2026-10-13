import openpyxl
import datetime
import hashlib
import re

wb = openpyxl.load_workbook('scratch_data.xlsx', data_only=True)

def sql_val(v):
    if v is None:
        return 'NULL'
    if isinstance(v, (datetime.datetime, datetime.date)):
        return f"'{v.strftime('%Y-%m-%d %H:%M:%S')}'"
    if isinstance(v, datetime.time):
        return f"'{v.strftime('%H:%M:%S')}'"
    if isinstance(v, bool):
        return '1' if v else '0'
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)
    s = str(v).strip()
    if s == '' or s.lower() == 'null':
        return 'NULL'
    escaped = s.replace('\\', '\\\\').replace("'", "''").replace('\r', '').replace('\n', '\\n')
    return f"'{escaped}'"

def int_or_null(v):
    if v is None:
        return 'NULL'
    try:
        return str(int(float(v)))
    except (ValueError, TypeError):
        return 'NULL'

sql_lines = [
    "-- ============================================================",
    "-- FULL DATABASE MIGRATION FROM GOOGLE SPREADSHEETS (OCT 2026)",
    "-- REPLACING ALL PREVIOUS DATA (CLEAN & ERROR-FREE)",
    "-- ============================================================",
    "USE db_maintenance_qr;",
    "SET NAMES utf8mb4;",
    "SET FOREIGN_KEY_CHECKS = 0;",
    "",
]

# 1. CABANG
ws = wb['Cabang']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 1. CABANG")
sql_lines.append("TRUNCATE TABLE cabang;")
for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        cid = int_or_null(r[0])
        if cid == 'NULL': continue
        nama = sql_val(r[1])
        alamat = sql_val(r[2] if len(r) > 2 else None)
        telp = sql_val(r[3] if len(r) > 3 else None)
        pj = sql_val(r[4] if len(r) > 4 else None)
        sql_lines.append(f"INSERT INTO cabang (id, nama_cabang, alamat, telepon, penanggung_jawab) VALUES ({cid}, {nama}, {alamat}, {telp}, {pj});")
sql_lines.append("")

# 2. DIVISI
ws = wb['Divisi']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 2. DIVISI")
sql_lines.append("TRUNCATE TABLE divisi;")
for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        did = int_or_null(r[0])
        if did == 'NULL': continue
        nama = sql_val(r[1])
        # id_cabang harus berupa integer atau NULL, bukan string teks seperti 'Internal Control Unit'
        id_cabang = int_or_null(r[2] if len(r) > 2 else None)
        sql_lines.append(f"INSERT INTO divisi (id, nama_divisi, id_cabang) VALUES ({did}, {nama}, {id_cabang});")
sql_lines.append("")

# 3. KATEGORI_ASET
ws = wb['Kategori_Aset']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 3. KATEGORI_ASET")
sql_lines.append("TRUNCATE TABLE kategori_aset;")
for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        kid = int_or_null(r[0])
        if kid == 'NULL': continue
        nama = sql_val(r[1])
        sql_lines.append(f"INSERT INTO kategori_aset (id, nama_kategori) VALUES ({kid}, {nama});")
sql_lines.append("")

# 4. KARYAWAN
ws = wb['Karyawan']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 4. KARYAWAN")
sql_lines.append("TRUNCATE TABLE karyawan;")
for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        kid = int_or_null(r[0])
        if kid == 'NULL': continue
        nama = sql_val(r[1])
        id_cab = int_or_null(r[2] if len(r) > 2 else 1)
        id_div = int_or_null(r[3] if len(r) > 3 else 1)
        sql_lines.append(f"INSERT INTO karyawan (id, nama_karyawan, id_cabang, id_divisi) VALUES ({kid}, {nama}, {id_cab}, {id_div});")
sql_lines.append("")

# 5. ASSETS
ws = wb['Assets']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 5. ASSETS")
sql_lines.append("TRUNCATE TABLE assets;")
valid_asset_ids = set()
for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        aid = int_or_null(r[0])
        if aid == 'NULL': continue
        valid_asset_ids.add(int(aid))
        kode = sql_val(r[1])
        merk = sql_val(r[2])
        model = sql_val(r[3])
        sn = sql_val(r[4])
        id_kat = int_or_null(r[5] if len(r) > 5 else 1)
        id_cab = int_or_null(r[6] if len(r) > 6 else 1)
        id_div = int_or_null(r[7] if len(r) > 7 else 1)
        id_kar = int_or_null(r[8] if len(r) > 8 else None)
        status = sql_val(r[9] if len(r) > 9 else 'Aktif')
        ket = sql_val(r[10] if len(r) > 10 else None)
        
        # IP and Printer
        ip_raw = r[11] if len(r) > 11 else None
        if ip_raw is not None:
            ip_clean = str(ip_raw).strip().replace(',', '.')
        else:
            ip_clean = None
        ip = sql_val(ip_clean)
        
        printer = sql_val(r[12] if len(r) > 12 else None)
        
        # Default created_at diatur ke 2024-01-01 agar cut-off filter periode bulan lalu (misal Sept 2026) membaca seluruh aset ini
        c_at_val = r[13] if len(r) > 13 and r[13] is not None else '2024-01-01 00:00:00'
        created_at = sql_val(c_at_val)
        
        sql_lines.append(f"INSERT INTO assets (id, kode_inventaris, merk, model, serial_number, id_kategori, id_cabang, id_divisi, id_karyawan, status, keterangan, ip_address, printer, created_at) VALUES ({aid}, {kode}, {merk}, {model}, {sn}, {id_kat}, {id_cab}, {id_div}, {id_kar}, {status}, {ket}, {ip}, {printer}, {created_at});")
sql_lines.append("")

# 6. ASSET_QR_TOKENS
ws = wb['Asset_QR_Tokens']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 6. ASSET_QR_TOKENS")
sql_lines.append("TRUNCATE TABLE asset_qr_tokens;")
covered_assets = set()
for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '' and r[1] is not None and str(r[1]).strip() != '':
        aid = int_or_null(r[1])
        if aid == 'NULL': continue
        tok = sql_val(r[2])
        label = sql_val(r[3] if len(r) > 3 else 'Bodi Casing')
        active = int_or_null(r[4] if len(r) > 4 else 1)
        c_at = sql_val(r[5] if len(r) > 5 and r[5] is not None else '2024-01-01 00:00:00')
        covered_assets.add(int(aid))
        sql_lines.append(f"INSERT INTO asset_qr_tokens (asset_id, token, placement_label, is_active, created_at) VALUES ({aid}, {tok}, {label}, {active}, {c_at});")

# Generate static tokens for remaining assets not in Asset_QR_Tokens
for aid in sorted(valid_asset_ids - covered_assets):
    tok = hashlib.sha256(f"STATIC_QR_MAINTENANCE_KEY_SALT_{aid}".encode()).hexdigest()[:32]
    sql_lines.append(f"INSERT INTO asset_qr_tokens (asset_id, token, placement_label, is_active, created_at) VALUES ({aid}, '{tok}', 'Bodi Casing', 1, '2024-01-01 00:00:00');")
sql_lines.append("")

# 7. USERS
ws = wb['Users']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 7. USERS")
sql_lines.append("TRUNCATE TABLE users;")
for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        uid = int_or_null(r[0])
        if uid == 'NULL': continue
        uname = sql_val(r[1])
        pw = sql_val(r[2])
        nama = sql_val(r[3])
        role = sql_val(r[4])
        telp = sql_val(r[5] if len(r) > 5 else None)
        stat = sql_val(r[6] if len(r) > 6 else 'Aktif')
        c_at = sql_val(r[7] if len(r) > 7 and r[7] is not None else datetime.datetime.now())
        face_desc = sql_val(r[8] if len(r) > 8 else None)
        face_photo = sql_val(r[9] if len(r) > 9 else None)
        face_status = sql_val(r[10] if len(r) > 10 else 'none')
        passkey = sql_val(r[11] if len(r) > 11 else None)
        panggilan = sql_val(r[12] if len(r) > 12 else None)
        sql_lines.append(f"INSERT INTO users (id, username, password, nama, role, telepon, status, created_at, face_descriptor, face_photo, face_status, passkey_credential, nama_panggilan) VALUES ({uid}, {uname}, {pw}, {nama}, {role}, {telp}, {stat}, {c_at}, {face_desc}, {face_photo}, {face_status}, {passkey}, {panggilan});")
sql_lines.append("")

# 8. MAINTENANCE_SCAN
ws = wb['Maintenance_Scan']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 8. MAINTENANCE_SCAN")
sql_lines.append("TRUNCATE TABLE maintenance_scan;")
for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        mid = int_or_null(r[0])
        aid = int_or_null(r[1])
        if mid == 'NULL' or aid == 'NULL': continue
        uid = int_or_null(r[2])
        tech_name = sql_val(r[3])
        
        # Tanggal & waktu maintenance
        m_date_val = r[4]
        if isinstance(m_date_val, (datetime.datetime, datetime.date)):
            m_date = f"'{m_date_val.strftime('%Y-%m-%d')}'"
        else:
            m_date = sql_val(m_date_val)
            
        m_time_val = r[5]
        if isinstance(m_time_val, datetime.time):
            m_time = f"'{m_time_val.strftime('%H:%M:%S')}'"
        else:
            m_time = sql_val(m_time_val)
            
        m_month = int_or_null(r[6])
        m_year = int_or_null(r[7])
        stat = sql_val(r[8] if len(r) > 8 else 'Selesai')
        source = sql_val(r[9] if len(r) > 9 else 'QR')
        c_at = sql_val(r[10] if len(r) > 10 and r[10] is not None else datetime.datetime.now())
        findings = sql_val(r[11] if len(r) > 11 else None)
        recommend = sql_val(r[12] if len(r) > 12 else None)
        bio_ver = int_or_null(r[13] if len(r) > 13 else 0)
        bio_conf = sql_val(r[14] if len(r) > 14 else None)
        bio_photo = sql_val(r[15] if len(r) > 15 else None)
        lat = sql_val(r[16] if len(r) > 16 else None)
        lng = sql_val(r[17] if len(r) > 17 else None)
        sql_lines.append(f"INSERT INTO maintenance_scan (id, asset_id, technician_user_id, technician_name, maintenance_date, maintenance_time, maintenance_month, maintenance_year, status, source, created_at, findings, recommendation, biometric_verified, biometric_confidence, biometric_photo, latitude, longitude) VALUES ({mid}, {aid}, {uid}, {tech_name}, {m_date}, {m_time}, {m_month}, {m_year}, {stat}, {source}, {c_at}, {findings}, {recommend}, {bio_ver}, {bio_conf}, {bio_photo}, {lat}, {lng});")
sql_lines.append("")

# 9. MAINTENANCE_CHECKLISTS
ws = wb['Maintenance_Checklists']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 9. MAINTENANCE_CHECKLISTS")
sql_lines.append("TRUNCATE TABLE maintenance_checklists;")
for r in rows:
    if r[0] is not None and str(r[0]).strip() != '':
        mid = int_or_null(r[1])
        aid = int_or_null(r[2])
        chk_num = int_or_null(r[3])
        if mid == 'NULL' or aid == 'NULL' or chk_num == 'NULL': continue
        chk_name = sql_val(r[4])
        checked = int_or_null(r[5] if len(r) > 5 else 1)
        notes = sql_val(r[6] if len(r) > 6 else None)
        c_at = sql_val(r[7] if len(r) > 7 and r[7] is not None else datetime.datetime.now())
        # Hindari memasukkan id sheet untuk mencegah error duplicate key (18 duplikat di spreadsheet)
        sql_lines.append(f"INSERT INTO maintenance_checklists (maintenance_id, asset_id, checklist_number, checklist_name, checked, notes, created_at) VALUES ({mid}, {aid}, {chk_num}, {chk_name}, {checked}, {notes}, {c_at});")
sql_lines.append("")

# 10. MAINTENANCE_FINDINGS
ws = wb['Maintenance_Findings']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 10. MAINTENANCE_FINDINGS")
sql_lines.append("DROP TABLE IF EXISTS maintenance_findings;")
sql_lines.append("""CREATE TABLE maintenance_findings (
    id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT PRIMARY KEY,
    maintenance_scan_id BIGINT UNSIGNED NULL,
    asset_id INT NOT NULL,
    created_by INT NULL,
    finding TEXT NOT NULL,
    action_taken TEXT NULL,
    severity VARCHAR(20) NOT NULL DEFAULT 'Sedang',
    repair_status VARCHAR(50) NOT NULL DEFAULT 'Open',
    status VARCHAR(50) NULL,
    reported_by VARCHAR(150) NULL,
    reported_at DATETIME NULL,
    resolved_by VARCHAR(150) NULL,
    resolved_at DATETIME NULL,
    catatan_penyelesaian TEXT NULL,
    created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
    KEY idx_finding_asset (asset_id),
    KEY idx_finding_status (repair_status)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;""")

for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        mid = int_or_null(r[1])
        aid = int_or_null(r[2])
        if aid == 'NULL': continue
        desc = sql_val(r[4] if len(r) > 4 else '-')
        action = sql_val(r[5] if len(r) > 5 else None)
        stat = sql_val(r[6] if len(r) > 6 else 'Open')
        rep_by = sql_val(r[7] if len(r) > 7 else None)
        rep_at = sql_val(r[8] if len(r) > 8 and r[8] is not None else datetime.datetime.now())
        res_by = sql_val(r[9] if len(r) > 9 else None)
        res_at = sql_val(r[10] if len(r) > 10 and r[10] is not None else None)
        notes = sql_val(r[11] if len(r) > 11 else None)
        sql_lines.append(f"INSERT INTO maintenance_findings (maintenance_scan_id, asset_id, finding, action_taken, severity, repair_status, status, reported_by, reported_at, resolved_by, resolved_at, catatan_penyelesaian, created_at) VALUES ({mid}, {aid}, {desc}, {action}, 'Sedang', {stat}, {stat}, {rep_by}, {rep_at}, {res_by}, {res_at}, {notes}, {rep_at});")
sql_lines.append("")

# 11. AUDIT_TRAIL -> system_audit_logs
ws = wb['Audit_Trail']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 11. AUDIT_TRAIL (system_audit_logs)")
sql_lines.append("TRUNCATE TABLE system_audit_logs;")
for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        c_at = sql_val(r[1] if len(r) > 1 and r[1] is not None else datetime.datetime.now())
        uid = int_or_null(r[2] if len(r) > 2 else None)
        uname = sql_val(r[3] if len(r) > 3 else 'Sistem')
        urole = sql_val(r[4] if len(r) > 4 else 'admin')
        ip = sql_val(r[5] if len(r) > 5 else '127.0.0.1')
        act = sql_val(r[6] if len(r) > 6 else 'UPDATE')
        mod = sql_val(r[7] if len(r) > 7 else 'ASET')
        tid = int_or_null(r[8] if len(r) > 8 else None)
        tlabel = sql_val(r[9] if len(r) > 9 else None)
        details = sql_val(r[10] if len(r) > 10 else None)
        # Hindari memasukkan fixed id untuk mencegah error duplikasi id di sheet (id 9, 22, 86, 99)
        sql_lines.append(f"INSERT INTO system_audit_logs (created_at, user_id, user_name, user_role, ip_address, action, module, target_id, target_label, details) VALUES ({c_at}, {uid}, {uname}, {urole}, {ip}, {act}, {mod}, {tid}, {tlabel}, {details});")
sql_lines.append("")

# 12. INVENTARIS_KARTU
ws = wb['inventaris_kartu']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 12. INVENTARIS_KARTU")
sql_lines.append("TRUNCATE TABLE inventaris_kartu;")
for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        norek = sql_val(r[1])
        nama = sql_val(r[2])
        tgl_val = r[3] if len(r) > 3 and r[3] is not None else datetime.date(2020, 1, 1)
        if isinstance(tgl_val, (datetime.datetime, datetime.date)):
            tgl = f"'{tgl_val.strftime('%Y-%m-%d')}'"
        else:
            tgl = sql_val(tgl_val)
        bc = sql_val(r[4] if len(r) > 4 else '')
        lok = sql_val(r[5] if len(r) > 5 else 'Kantor Pusat (KPO)')
        c_at = sql_val(r[6] if len(r) > 6 and r[6] is not None else datetime.datetime.now())
        # Hindari duplikat id (id 9, 10, 11, 12, 13 duplikat di spreadsheet)
        sql_lines.append(f"INSERT INTO inventaris_kartu (nomor_rekening, nama_barang, tanggal_perolehan, barcode_data, lokasi, pengguna, created_at) VALUES ({norek}, {nama}, {tgl}, {bc}, {lok}, 'Umum / Pool', {c_at});")
sql_lines.append("")

# 13. EMPLOYEE_COMPLAINTS
ws = wb['Employee_Complaints']
rows = list(ws.iter_rows(values_only=True))
sql_lines.append("-- 13. EMPLOYEE_COMPLAINTS")
sql_lines.append("DROP TABLE IF EXISTS employee_complaints;")
sql_lines.append("""CREATE TABLE employee_complaints (
    id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT PRIMARY KEY,
    asset_id INT NOT NULL,
    reporter_name VARCHAR(150) NOT NULL,
    kategori VARCHAR(50) NOT NULL DEFAULT 'Lain-lain',
    complaint TEXT NOT NULL,
    contact VARCHAR(50) NULL,
    status VARCHAR(50) NOT NULL DEFAULT 'Menunggu Teknisi',
    technician_response TEXT NULL,
    resolved_by VARCHAR(150) NULL,
    resolved_at DATETIME NULL,
    created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
    KEY idx_asset (asset_id),
    KEY idx_status (status)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;""")

for r in rows[1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        aid = int_or_null(r[1])
        if aid == 'NULL': continue
        emp_name = sql_val(r[2])
        kat = sql_val(r[3] if len(r) > 3 else 'Umum')
        comp_text = sql_val(r[4])
        contact = sql_val(r[5] if len(r) > 5 else None)
        stat = sql_val(r[6] if len(r) > 6 else 'Menunggu Teknisi')
        tech_note = sql_val(r[7] if len(r) > 7 else None)
        res_by = sql_val(r[8] if len(r) > 8 else None)
        res_at = sql_val(r[9] if len(r) > 9 and r[9] is not None else None)
        c_at = sql_val(r[10] if len(r) > 10 and r[10] is not None else datetime.datetime.now())
        sql_lines.append(f"INSERT INTO employee_complaints (asset_id, reporter_name, kategori, complaint, contact, status, technician_response, resolved_by, resolved_at, created_at) VALUES ({aid}, {emp_name}, {kat}, {comp_text}, {contact}, {stat}, {tech_note}, {res_by}, {res_at}, {c_at});")
sql_lines.append("")

sql_lines.append("SET FOREIGN_KEY_CHECKS = 1;")

out_file = 'sql/05_full_database_migration.sql'
with open(out_file, 'w', encoding='utf-8') as f:
    f.write('\n'.join(sql_lines) + '\n')

print(f"Generated {out_file} with {len(sql_lines)} lines of SQL successfully!")
